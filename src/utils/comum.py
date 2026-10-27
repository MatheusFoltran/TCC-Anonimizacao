# Utilitários compartilhados por nucleo.py e pelos módulos de ataque/modelo.

from pathlib import Path
import json

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


def normalizar_colunas(df: pd.DataFrame, colunas: list[str]) -> pd.DataFrame:
    for col in colunas:
        if col in df.columns:
            serie = df[col].astype("string").str.strip()
            serie = serie.str.replace("Â°", "°", regex=False)
            df[col] = serie.replace(["", "nan", "None", "<NA>"], pd.NA)
    return df


def montar_chave(df: pd.DataFrame, colunas: list[str]) -> pd.Series:
    # Concatena as colunas com "|"; nulos viram "MISSING".
    partes = [df[col].fillna("MISSING").astype(str) for col in colunas]
    return partes[0].str.cat(partes[1:], sep="|")


QI_COLUNAS = ["DT_NASC", "CO_MUN_RES", "CS_SEXO", "CS_RACA", "CO_MUN_NOT"]
SENSIVEL_COLUNAS = ["EVOLUCAO", "CLASSI_FIN"]

# Seed fixa para amostragens aleatórias de alvos (homogeneidade, assimetria),
# para reprodutibilidade.
RANDOM_SEED = 10


def normalizar_qi_sensivel(df: pd.DataFrame) -> pd.DataFrame:
    colunas = [c for c in QI_COLUNAS + SENSIVEL_COLUNAS if c in df.columns]
    return normalizar_colunas(df, colunas)


def montar_chave_qi(df: pd.DataFrame) -> pd.Series:
    return montar_chave(df, QI_COLUNAS)


def outro_atributo_sensivel(sensivel_alvo: str) -> str:
    outros = [c for c in SENSIVEL_COLUNAS if c != sensivel_alvo]
    return outros[0]


def montar_chave_bloco(df: pd.DataFrame, sensivel_alvo: str) -> pd.Series:
    # Bloco multi-atributo (l_diversidade/assimetria): QI + o outro sensível.
    return montar_chave(df, QI_COLUNAS + [outro_atributo_sensivel(sensivel_alvo)])


def razao_segura(numerador, denominador, casas: int = 4) -> float:
    # 0.0 (em vez de ZeroDivisionError) quando denominador e' 0.
    if not denominador:
        return 0.0
    return round(numerador / denominador, casas)


def percentual(numerador, denominador, casas: int = 4) -> float:
    return razao_segura(numerador * 100, denominador, casas)


def round_seguro(valor, tamanho, casas: int = 4) -> float:
    # 0.0 (em vez de NaN) quando a serie de origem de "valor" esta vazia.
    if not tamanho:
        return 0.0
    return round(float(valor), casas)


def validar_colunas(df: pd.DataFrame, sensivel_col: str | None = None) -> None:
    if sensivel_col is not None and sensivel_col not in df.columns:
        raise ValueError(f"Coluna {sensivel_col} nao encontrada na base de entrada.")
    for col in QI_COLUNAS:
        if col not in df.columns:
            raise ValueError(f"Coluna obrigatoria ausente: {col}")


def get_input_path() -> Path:
    # candidatos.csv: entrada padrão gerada por k_anonimato.py.
    return Path(__file__).resolve().parent.parent / "modelos" / "k_anonimato" / "resultados" / "candidatos.csv"


def caminho_relativo(path: Path) -> str:
    # Caminho relativo à raiz do projeto, com "/" em qualquer SO, para que os
    # resumos JSON não registrem o caminho absoluto da máquina que os gerou.
    raiz = Path(__file__).resolve().parent.parent.parent
    return Path(path).resolve().relative_to(raiz).as_posix()


def carregar_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(f"Arquivo nao encontrado: {path}")
    try:
        df = pd.read_csv(path, sep=";", dtype=str, low_memory=False, encoding="utf-8")
    except UnicodeDecodeError:
        df = pd.read_csv(path, sep=";", dtype=str, low_memory=False, encoding="latin1")
    df.columns = [c.strip().replace('"', "") for c in df.columns]
    return df


def exportar_xlsx(df: pd.DataFrame, path: Path) -> None:
    # Linhas zebradas + largura de coluna ajustada ao maior valor.
    wb = Workbook()
    ws = wb.active
    fill_a = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
    fill_b = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

    ws.append(list(df.columns))
    for idx, row in enumerate(df.itertuples(index=False), start=1):
        valores = []
        for v in row:
            if pd.isna(v):
                valores.append(None)
            else:
                valores.append(v)
        ws.append(valores)

        if idx % 2:
            fill = fill_a
        else:
            fill = fill_b
        for col in range(1, len(df.columns) + 1):
            ws.cell(row=idx + 1, column=col).fill = fill

    for col_idx, coluna in enumerate(df.columns, start=1):
        letra = get_column_letter(col_idx)
        valores = [str(coluna)]
        for v in df[coluna].tolist():
            if pd.isna(v):
                valores.append("")
            else:
                valores.append(str(v))
        ws.column_dimensions[letra].width = min(max(len(v) for v in valores) + 2, 100)

    wb.save(path)


def salvar_resumo_json(resumo: dict, path: Path) -> None:
    with path.open("w", encoding="utf-8") as f:
        json.dump(resumo, f, ensure_ascii=False, indent=2)
